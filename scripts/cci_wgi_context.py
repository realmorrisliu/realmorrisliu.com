"""Extract WGI aggregates and link country context, never city governance scores."""

import argparse
import csv
import hashlib
import json
import math
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SOURCE_HASH = "67c38d2e9253fb1a72ee302c3b7c47b2ba7b177a04673163796ca1666f506ecb"
DIMENSIONS = ("va", "pv", "ge", "rq", "rl", "cc")
HEADERS = (
    "ID variable (economy code/ gov. dimension/ year)", "Economy (name)", "Economy (code)",
    "Region", "Income classification", "Year", "Governance dimension", "Number of sources",
    "Governance estimate (approx. -2.5 to +2.5)", "Standard error (estimate)",
    "Lower threshold (90% conf. int. estimate)", "Upper threshold (90% conf. int. estimate)",
    "Governance score (0-100)", "Standard error (gov. score)",
    "Lower threshold (90% conf. int. score)", "Upper threshold (90% conf. int. score)",
)
FIELDS = ("source_id", "economy_name", "economy_code", "region", "income_classification", "year",
          "dimension", "number_of_sources", "estimate", "estimate_standard_error", "estimate_lower90",
          "estimate_upper90", "score", "score_standard_error", "score_lower90", "score_upper90")


def aggregate(row, dimension):
    if len(row) != 16 or row[5] != 2025 or row[6] != dimension or row[0] != f"{row[2]}{dimension}2025":
        raise ValueError("Unexpected aggregate identity")
    if not isinstance(row[7], int) or row[7] < 1:
        raise ValueError("Invalid source count")
    if any(not isinstance(v, (int, float)) or not math.isfinite(v) for v in row[8:]):
        raise ValueError("Missing or nonfinite aggregate")
    if not (row[10] <= row[8] <= row[11] and 0 <= row[14] <= row[12] <= row[15] <= 100):
        raise ValueError("Invalid aggregate interval")
    if row[9] < 0 or row[13] < 0:
        raise ValueError("Negative standard error")
    return dict(zip(FIELDS, row))


def country_links(candidates, aggregates):
    available = {r["source_id"] for r in aggregates}
    if len(available) != len(aggregates):
        raise ValueError("Duplicate aggregate identity")
    results = []
    for candidate in candidates:
        code = candidate["country_iso"]
        mapped = "XKX" if code == "XKO" else code
        refs = {dimension + "_country_context_ref": f"{mapped}{dimension}2025"
                if f"{mapped}{dimension}2025" in available else "" for dimension in DIMENSIONS}
        results.append({"efua_id": candidate["efua_id"], "source_name": candidate["source_name"],
                        "fua_country_code": code, "lookup_economy_code": mapped,
                        "code_mapping": "verified_kosovo_alias" if code == "XKO" else "exact_code",
                        "evidence_scale": "country_context_not_city_measurement",
                        "available_dimensions": sum(bool(v) for v in refs.values()), **refs})
    return results


def main():
    import openpyxl

    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", required=True, type=Path)
    args = parser.parse_args()
    if hashlib.sha256(args.xlsx.read_bytes()).hexdigest() != SOURCE_HASH:
        raise ValueError("Workbook checksum mismatch")
    folder = ROOT / "docs/data"
    candidate_path = folder / "cci-global-candidates-r2019a.csv"
    provenance = json.loads((folder / "cci-global-candidates-r2019a.provenance.json").read_text())
    if hashlib.sha256(candidate_path.read_bytes()).hexdigest() != provenance["csvSha256"]:
        raise ValueError("Candidate checksum mismatch")
    with candidate_path.open() as stream:
        candidates = list(csv.DictReader(stream))
    if len(candidates) != provenance["rowCount"] or len({r["efua_id"] for r in candidates}) != len(candidates):
        raise ValueError("Candidate count or identity mismatch")
    workbook = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    if tuple(workbook.sheetnames) != DIMENSIONS:
        raise ValueError("Unexpected workbook sheets")
    aggregates = []
    for dimension in DIMENSIONS:
        # Only the aggregate columns: third-party source values are not exported.
        rows = workbook[dimension].iter_rows(max_col=16, values_only=True)
        if tuple(next(rows)) != HEADERS:
            raise ValueError("Workbook schema changed")
        for row in rows:
            if row[5] == 2025:
                aggregates.append(aggregate(row, dimension))
    workbook.close()
    links = country_links(candidates, aggregates)
    outputs = {"cci-wgi-2025-country-aggregates.csv": aggregates, "cci-wgi-efua-context.csv": links}
    output_hashes = {}
    for name, rows in outputs.items():
        path = folder / name
        with path.open("w", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]), lineterminator="\n")
            writer.writeheader()
            writer.writerows(rows)
        output_hashes[name] = hashlib.sha256(path.read_bytes()).hexdigest()
    manifest = {
        "source": "https://www.worldbank.org/content/dam/sites/govindicators/doc/wgidataset_with_sourcedata-2026.xlsx",
        "citation": "Worldwide Governance Indicators, 2026 Update, World Bank; accessed 2026-10-10; extracted 2025 aggregates only.",
        "releaseYear": 2026, "observationYear": 2025, "workbookSha256": SOURCE_HASH,
        "candidateCsvSha256": provenance["csvSha256"], "candidateCount": len(candidates),
        "countryAggregateCountByDimension": dict(Counter(r["dimension"] for r in aggregates)),
        "candidateAvailableDimensionCounts": dict(Counter(r["available_dimensions"] for r in links)),
        "countryCodeAlias": {"XKO": {"target": "XKX", "evidence": "Frozen FUA Cntry_name=Kosovo; WGI Economy (name)=Kosovo."}},
        "missingCountryCodesByDimension": {d: sorted({r["fua_country_code"] for r in links if not r[d + "_country_context_ref"]}) for d in DIMENSIONS},
        "scope": "Country context only. No aggregate is repeated as a city measurement; no CCI scores.",
        "excludedColumns": "All source-level means (columns 17 onward); workbook remains outside repository.",
        "licence": "World Bank CC BY 4.0 with applicable additional terms; third-party rescaled inputs excluded.",
        "licencePolicy": "https://datacatalog.worldbank.org/public-licenses",
        "usageAdvisory": "https://www.worldbank.org/en/publication/worldwide-governance-indicators/usage-advisory",
        "scriptSha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "openpyxlVersion": openpyxl.__version__, "outputs": output_hashes,
    }
    (folder / "cci-wgi-context.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(manifest["countryAggregateCountByDimension"], manifest["candidateAvailableDimensionCounts"])


if __name__ == "__main__":
    main()
